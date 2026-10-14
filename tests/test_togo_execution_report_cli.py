"""scripts/togo_execution_report.py writes one year's tables from the database."""
import importlib.util
import pathlib
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from queries import QueryService
from tests.execution_fixture import TOTAL_FINALE, YEAR, execution_frame

SCRIPT = pathlib.Path(__file__).resolve().parent.parent / "scripts" / "togo_execution_report.py"


def _load_script():
    spec = importlib.util.spec_from_file_location("togo_execution_report_cli", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class CliTest(unittest.TestCase):

    def tearDown(self):
        QueryService._instance = None

    @patch("queries.PUBLIC_ONLY", False)
    def test_writes_the_year_from_a_fresh_read(self):
        cli = _load_script()
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(QueryService, "get_togo_execution_data",
                             return_value=execution_frame()) as read:
            out = pathlib.Path(tmp) / "report.xlsx"
            self.assertEqual(cli.main(["--year", str(YEAR), "--out", str(out)]), 0)
            ws = load_workbook(out)["Natures"]
        read.assert_called_once_with(persistent=False)
        self.assertEqual(ws.cell(ws.max_row, 5).value, TOTAL_FINALE)

    @patch("queries.PUBLIC_ONLY", False)
    def test_year_without_lines_fails(self):
        cli = _load_script()
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(QueryService, "get_togo_execution_data", return_value=execution_frame()):
            out = pathlib.Path(tmp) / "report.xlsx"
            self.assertEqual(cli.main(["--year", "1990", "--out", str(out)]), 1)
            self.assertFalse(out.exists())


if __name__ == "__main__":
    unittest.main()
