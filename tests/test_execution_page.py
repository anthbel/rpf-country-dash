"""Callbacks of the « Exécution du budget » page."""
import base64
import io
import unittest
from unittest.mock import patch

from dash.exceptions import PreventUpdate
from openpyxl import load_workbook

import server_store
from components import togo_execution as te
from translations import t
from tests.execution_fixture import PRIOR_YEAR, YEAR, execution_frame

with patch("dash.register_page"):
    from pages import execution as page


def _missing():
    raise RuntimeError('relation "boost.tgo_execution" does not exist')


class PageTest(unittest.TestCase):

    def setUp(self):
        server_store.clear()
        self.addCleanup(server_store.clear)

    def _load_fixture(self):
        server_store.set("togo_execution", te.normalise_execution(execution_frame()))

    def test_missing_table_shows_unavailable(self):
        with patch.dict("data_mapping.function_data_mapping", {"togo_execution": _missing}):
            store = page.fetch_execution_data("Togo")
            self.assertEqual(store, {"country": "Togo", "ready": False})
            self.assertEqual(page.update_year_options(store, "Togo", None), ([], None, True))
            fig, text = page.update_overview(store, "Togo", None, "fr")
        self.assertEqual(text, t("error.data_unavailable", "fr"))
        self.assertFalse(fig.data)

    def test_other_country_shows_unavailable(self):
        self._load_fixture()
        store = page.fetch_execution_data("Kenya")
        self.assertFalse(store["ready"])
        _, text = page.update_stages(store, "Kenya", None, "fr")
        self.assertEqual(text, t("error.data_unavailable", "fr"))

    def test_years_default_to_the_latest(self):
        self._load_fixture()
        store = page.fetch_execution_data("Togo")
        options, value, disabled = page.update_year_options(store, "Togo", None)
        self.assertEqual([o["value"] for o in options], [YEAR, PRIOR_YEAR])
        self.assertEqual(value, YEAR)
        self.assertFalse(disabled)
        self.assertEqual(page.update_year_options(store, "Togo", PRIOR_YEAR)[1], PRIOR_YEAR)

    def test_stale_store_waits(self):
        with self.assertRaises(PreventUpdate):
            page.update_overview({"country": "Kenya", "ready": False}, "Togo", YEAR, "fr")

    def test_blocks_and_tables_render(self):
        self._load_fixture()
        store = page.fetch_execution_data("Togo")
        for update in (page.update_overview, page.update_stages, page.update_ministries,
                       page.update_ministry_nature, page.update_investment, *page.update_tags.values()):
            fig, text = update(store, "Togo", YEAR, "fr")
            self.assertTrue(fig.data)
            self.assertNotEqual(text, t("error.data_unavailable", "fr"))
        ministries, projects = page.update_tables(store, "Togo", YEAR, "fr")
        self.assertIsNotNone(ministries)
        self.assertIsNotNone(projects)
        options, value = page.update_ministry_options(store, "Togo", YEAR, "fr", None)
        self.assertEqual(value, "210")
        table, text = page.update_programmes(store, "Togo", YEAR, "320", "fr")
        self.assertIsNotNone(table)
        self.assertIn("MINISTERE DE LA SANTE", text)

    def test_download_is_the_year_report(self):
        self._load_fixture()
        data = page.download_report(1, YEAR, "Togo")
        self.assertEqual(data["filename"], f"execution_budget_togo_{YEAR}.xlsx")
        wb = load_workbook(io.BytesIO(base64.b64decode(data["content"])))
        self.assertEqual(wb.sheetnames[0], "Natures")


if __name__ == "__main__":
    unittest.main()
